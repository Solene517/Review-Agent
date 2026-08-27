import os
import re
import time
import json
from openai import OpenAI
from dotenv import load_dotenv
from pptx import Presentation
from pypdf import PdfReader

load_dotenv()
client = OpenAI(
    api_key=os.getenv("LLM_API_KEY"),
    base_url=os.getenv("LLM_BASE_URL")
)
model = os.getenv("LLM_MODEL", "deepseek-chat")

def call_llm(messages, temperature=0.3, max_tokens=5000):
    response = client.chat.completions.create(
        model=model,
        messages=messages,
        temperature=temperature,
        max_tokens=max_tokens,
    )
    return response.choices[0].message.content

def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_runs.append(shape.text.strip())
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_runs.append(cell.text.strip())
    return "\n".join(text_runs)

def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    text_runs = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_runs.append(text.strip())
    return "\n".join(text_runs)

def extract_topics(chapter_text):
    prompt = f"""
你是一位资深教师，请阅读以下章节内容，提取本章节的核心考点。

要求：
1. 只提取学科知识点，不要提取“常见题型”、“应用”等非知识点条目。
2. 考点数量控制在 3~4 个，只保留最核心、最重要的知识点。
3. 每个考点应覆盖一个核心概念或原理。
4. 输出格式：每行一个考点，格式为“编号. 中文名称 / English Name - 一句话描述”，例如：
1. 高斯定律 / Gauss's Law - 描述电场与电荷的关系
5. 直接输出列表，不要添加任何其他文字或代码块标记。

章节内容：
{chapter_text}
"""
    result = call_llm([{"role": "user", "content": prompt}], max_tokens=1000)
    topics = []
    lines = result.strip().split('\n')
    for line in lines:
        line = line.strip()
        if not line or not re.match(r'^\d+\.', line):
            continue
        match = re.match(r'^\d+\.\s*(.*?)\s*/\s*(.*?)\s*-\s*(.*)$', line)
        if match:
            zh = match.group(1).strip()
            en = match.group(2).strip()
            brief = match.group(3).strip()
            topics.append({"id": len(topics)+1, "zh": zh, "en": en, "brief": brief})
        else:
            parts = line.split(' - ', 1)
            name_part = parts[0].strip()
            brief = parts[1].strip() if len(parts) > 1 else ""
            name_part = re.sub(r'^\d+\.\s*', '', name_part)
            if '/' in name_part:
                zh, en = name_part.split('/', 1)
                topics.append({"id": len(topics)+1, "zh": zh.strip(), "en": en.strip(), "brief": brief})
            else:
                topics.append({"id": len(topics)+1, "zh": name_part, "en": "", "brief": brief})
    topics = [t for t in topics if not any(kw in t['zh'] for kw in ['题型', '问题', '应用'])]
    if len(topics) > 4:
        topics = topics[:4]
    return topics

def generate_topic_details(topic, chapter_summary, exam_config, all_topics):
    # 动态构建例题要求（使用明确的占位说明，禁止省略号）
    example_lines = []

    for i in range(exam_config.get("multiple_choice", 0)):
        difficulty = "简单" if i == 0 else "中等"
        example_lines.append(f"选择题{i+1}（{difficulty}）")
        example_lines.append("中文：请写出完整的中文题目，不要使用省略号")
        example_lines.append("English：Please write the complete English question, no ellipsis")
        example_lines.append("A. 请写出选项A完整内容  B. 请写出选项B完整内容  C. 请写出选项C完整内容  D. 请写出选项D完整内容")
        example_lines.append("答案：请填写正确选项字母（A/B/C/D）")
        example_lines.append("解析：请写出详细解析，解释为什么选这个答案")
        example_lines.append("")

    for i in range(exam_config.get("fill_blank", 0)):
        example_lines.append(f"填空题{i+1}（简单/中等）")
        example_lines.append("中文：请写出完整的中文题目，不要使用省略号")
        example_lines.append("English：Please write the complete English question, no ellipsis")
        example_lines.append("答案：请写出标准答案")
        example_lines.append("")

    for i in range(exam_config.get("short_answer", 0)):
        difficulty = "中等" if i == 0 else "困难"
        example_lines.append(f"简答题{i+1}（{difficulty}）")
        example_lines.append("中文：请写出完整的中文题目，不要使用省略号")
        example_lines.append("English：Please write the complete English question, no ellipsis")
        example_lines.append("参考答案（中文）：请写出完整答案")
        example_lines.append("参考答案（English）：Please write the complete answer")
        example_lines.append("")

    for i in range(exam_config.get("calculation", 0)):
        difficulty = "中等" if i == 0 else "困难"
        example_lines.append(f"计算题大题{i+1}（{difficulty}）")
        example_lines.append("中文：请写出完整的中文题目，不要使用省略号")
        example_lines.append("English：Please write the complete English question, no ellipsis")
        example_lines.append("解题步骤（中文）：请写出详细步骤，分步列出")
        example_lines.append("解题步骤（English）：Please write complete steps")
        example_lines.append("最终答案：请写出最终结果")
        example_lines.append("")

    example_req_text = "\n".join(example_lines)

    # 构建其他考点简要列表，用于串联讲解
    other_topics = [t for t in all_topics if t['id'] != topic['id']]
    other_topics_str = "\n".join([f"{t['id']}. {t['zh']} / {t['en']}" for t in other_topics])

    prompt = f"""
你是一位资深教师，请针对以下考点生成一份**面向应试**的复习卡片，要求内容全面、结构清晰、中英双语，并包含详细的讲解和知识点串联。

考点名称：{topic['zh']} / {topic['en']}
章节上下文摘要：{chapter_summary}
本章其他考点（用于串联讲解）：
{other_topics_str if other_topics_str else "无"}

请严格按照以下 Markdown 格式输出：

### {topic['zh']} / {topic['en']}

**知识点详解**  
中文：  
（用3-5句话详细阐述该考点的核心概念、原理或推导思路，要清晰易懂）

English:  
(3-5 sentences explaining the core concept, principle, or derivation in clear English)

**通俗讲解与串联**  
中文：  
（用日常语言、类比或生活例子解释难点；然后说明该考点与本章其他考点的逻辑关系，例如“理解了A之后，才能理解B，因为……”）

English:  
(Explain the difficult parts using plain language, analogies, or real-life examples; then explain how this topic connects to other topics in this chapter.)

**中文关键点**  
- 关键点1：一句话解释这个关键点是什么或为什么重要  
- 关键点2：一句话解释  
（列出3-5个核心要点，每个要点后必须用冒号隔开并加上解释）

**English Key Points**  
- Key point 1: one-sentence explanation  
- Key point 2: one-sentence explanation  
（对应中文关键点，每个要点后必须加英文解释）

**公式与易错点**  
| 公式 | 条件 | 用途 |
|------|------|------|
| $公式1$ | 适用条件 | 主要用途 |
（至少列出1个重要公式，公式使用 LaTeX 格式，例如 $E_0 = cB_0$）

易错点：  
- 易错点1（具体例子：写出学生常犯的具体错误，例如“学生常把 X 与 Y 混淆”）  
- 易错点2（具体例子：写出另一个具体错误）  
（每个易错点后必须用括号给出一个具体、完整的错误例子，严禁使用省略号）

**应试技巧**  
- 中文技巧1（举例说明应用场景）  
- 中文技巧2（举例说明应用场景）  

**Exam Tips**  
- English tip 1 (with example)  
- English tip 2 (with example)  

**例题**  

{example_req_text}

要求：
1. 知识点详解和通俗讲解要详细、深入，不要只列要点。
2. 通俗讲解必须包含与其他考点的串联逻辑。
3. 例题必须紧扣考点，难度搭配合理。
4. 公式必须使用 LaTeX 格式（用 $ 包围），例如 $E_0 = cB_0$。
5. 所有中英文内容准确、地道。
6. 只输出上述格式的内容，不要添加额外解释或代码块。
7. 关键点不能只是名词，必须每个后面都有解释（中英文对应）。
8. 严禁在任何地方使用省略号（...）或“略”等占位符，所有内容必须是完整的、可直接使用的文字。
"""
    return call_llm([{"role": "user", "content": prompt}], max_tokens=6000)

def generate_mindmap(topics):
    """
    生成 Markdown 无序列表形式的知识框架，避免使用特殊符号。
    """
    topic_list = "\n".join([f"{t['id']}. {t['zh']} / {t['en']} - {t['brief']}" for t in topics])
    prompt = f"""
根据以下考点列表，生成一个**层次清晰的 Markdown 无序列表**，展示本章节知识点之间的逻辑关系。

要求：
1. 使用 Markdown 无序列表（用 - 或 *），通过缩进表示层级（每个层级缩进两个空格）。
2. 每个考点作为一级节点，下面列出 2-4 个子要点（例如关键概念、公式、物理意义等），子要点下可再列细项。
3. 节点内容使用中英文双语，例如“高斯定律 Gauss's Law”。
4. 不要使用 ├──、│、└── 等特殊符号，只使用普通的列表符号和缩进。
5. 直接输出 Markdown 列表，不要包含其他文字。

考点列表：
{topic_list}
"""
    return call_llm([{"role": "user", "content": prompt}], max_tokens=3000)

def generate_connection_text(topics):
    """
    根据考点列表，生成一段文字描述各个考点之间的逻辑串联关系。
    """
    topic_list = "\n".join([f"{t['id']}. {t['zh']} / {t['en']} - {t['brief']}" for t in topics])
    prompt = f"""
你是一位资深教师，请根据以下考点列表，写一段“知识点串联”文字，说明这些考点之间的逻辑关系和发展脉络。

要求：
1. 用中文写，语言流畅，适合学生阅读。
2. 解释为什么这些考点按此顺序排列，它们如何逐步深入或相互关联。
3. 可以使用“首先...然后...进一步...最终...”等过渡词。
4. 内容要基于考点之间的学科逻辑，不要凭空编造。
5. 输出纯文本段落，不要使用Markdown标记。

考点列表：
{topic_list}
"""
    return call_llm([{"role": "user", "content": prompt}], max_tokens=1500)

def generate_formula_summary(chapter_text, topics):
    topic_list = "\n".join([f"{t['id']}. {t['zh']} / {t['en']}" for t in topics])
    prompt = rf"""
你是一位资深教师，请根据以下章节内容和考点列表，生成一个**完整的公式总结**，必须覆盖本章所有重要公式，不能遗漏任何一个。

要求：
1. 先仔细阅读章节内容，找出所有公式（包括变体、推导公式、重要结论等）。
2. 使用**无序列表**形式，每个公式一行，格式如下：
   - 公式：$E=mc^2$ 名称：质能方程 条件/备注：任何情况
3. 公式必须使用 LaTeX 格式（用 $ 包围），且不要包含竖线 | 等特殊字符，如需表示绝对值或范数，请使用 \vert 或 \mid。
4. 每个考点下涉及的公式都要列出，如果考点较多，可以按考点分组，使用缩进或子标题。
5. 名称/含义用中文，条件备注用简洁中文。
6. 直接输出列表，不要添加任何其他文字。

章节内容摘要：
{chapter_text[:2000]}

考点列表：
{topic_list}
"""
    return call_llm([{"role": "user", "content": prompt}], max_tokens=4000)  # 增加到4000

def generate_review(chapter_text, exam_config):
    print("正在提取核心考点...")
    topics = extract_topics(chapter_text)
    if not topics:
        print("未能提取到考点，请检查输入或重试。")
        return None
    print(f"共提取到 {len(topics)} 个考点：")
    for t in topics:
        print(f"  {t['id']}. {t['zh']} / {t['en']}")

    chapter_summary = chapter_text[:600]
    all_details = []
    for topic in topics:
        print(f"正在生成考点 {topic['id']} 的复习卡片...")
        detail_text = generate_topic_details(topic, chapter_summary, exam_config, topics)
        if detail_text:
            all_details.append(detail_text)
            print("  生成成功。")
        else:
            print(f"  考点 {topic['id']} 生成失败，跳过")
        time.sleep(1)

    print("正在生成知识框架...")
    mindmap_text = generate_mindmap(topics)

    print("正在生成知识点串联...")
    connection_text = generate_connection_text(topics)

    print("正在生成公式总结...")
    formula_summary = generate_formula_summary(chapter_text, topics)

    # 组合：知识框架（Markdown列表）+ 串联 + 公式总结
    combined_framework = (
        "## 知识框架\n" + mindmap_text + "\n\n"
        "## 知识点串联\n" + connection_text + "\n\n"
        "## 公式总结\n" + formula_summary
    )

    result = {
        "topics": topics,
        "details": all_details,
        "mindmap": combined_framework,
        "formula_summary": formula_summary
    }

    with open("review_output.json", "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    md = []
    md.append("# 复习资料\n")
    md.append(combined_framework)
    md.append("\n## 核心考点复习卡片\n")
    for detail in all_details:
        md.append(detail)
        md.append("\n---\n")
    with open("review_output.md", "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    print("完成！输出文件：review_output.json 和 review_output.md")
    return result

if __name__ == "__main__":
    # 命令行入口（与之前相同，支持粘贴、txt、PPT、PDF）
    print("请选择章节内容输入方式：")
    print("1. 直接在终端粘贴（适合短章节）")
    print("2. 从 chapter.txt 文件读取（适合长章节，推荐）")
    print("3. 从 PPT 文件读取（.pptx）")
    print("4. 从 PDF 文件读取（.pdf）")
    choice = input("请输入 1、2、3 或 4：").strip()

    if choice == "1":
        print("\n请在下方粘贴章节内容（粘贴后按回车，然后输入 END 并回车结束）：")
        lines = []
        while True:
            line = input()
            if line.strip() == "END":
                break
            lines.append(line)
        chapter_text = "\n".join(lines)
        if not chapter_text.strip():
            print("未输入任何内容，使用默认示例。")
            chapter_text = "麦克斯韦方程组是描述电场、磁场与电荷密度、电流密度之间关系的四个基本方程。"
    elif choice == "2":
        file_path = "chapter.txt"
        if not os.path.exists(file_path):
            print(f"未找到 {file_path}，请先创建该文件并写入章节内容。")
            exit(1)
        with open(file_path, "r", encoding="utf-8") as f:
            chapter_text = f.read()
        if not chapter_text.strip():
            print("chapter.txt 文件为空，请写入内容后重试。")
            exit(1)
    elif choice == "3":
        file_path = input("请输入 PPT 文件的路径（例如：C:/Users/.../复习.pptx）：").strip()
        file_path = file_path.strip('"').strip("'")
        if not os.path.exists(file_path):
            print(f"未找到文件：{file_path}")
            exit(1)
        print("正在提取 PPT 文本...")
        chapter_text = extract_text_from_pptx(file_path)
        if not chapter_text.strip():
            print("PPT 文件中没有提取到文本，请确认内容。")
            exit(1)
        print("PPT 文本提取成功，开始生成复习资料...")
    elif choice == "4":
        file_path = input("请输入 PDF 文件的路径（例如：C:/Users/.../复习.pdf）：").strip()
        file_path = file_path.strip('"').strip("'")
        if not os.path.exists(file_path):
            print(f"未找到文件：{file_path}")
            exit(1)
        print("正在提取 PDF 文本...")
        chapter_text = extract_text_from_pdf(file_path)
        if not chapter_text.strip():
            print("PDF 文件中没有提取到文本，可能是扫描版或图片型 PDF，请尝试其他方式。")
            exit(1)
        print("PDF 文本提取成功，开始生成复习资料...")
    else:
        print("输入无效，退出。")
        exit(1)

    print("\n请设置每种题型的数量（输入数字，0 表示不要该题型）：")
    try:
        num_choice = int(input("选择题数量（建议2-3）："))
        num_fill = int(input("填空题数量（建议0-1）："))
        num_short = int(input("简答题数量（建议1-2）："))
        num_calc = int(input("计算题大题数量（建议0-1）："))
    except ValueError:
        print("输入无效，使用默认值：选择题2，填空题0，简答题2，计算题0")
        num_choice, num_fill, num_short, num_calc = 2, 0, 2, 0

    exam_config = {
        "multiple_choice": num_choice,
        "fill_blank": num_fill,
        "short_answer": num_short,
        "calculation": num_calc
    }

    generate_review(chapter_text, exam_config)